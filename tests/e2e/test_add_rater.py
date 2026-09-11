from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from metaphor_agreement_studio.domain.imports import ValidationDecision
from metaphor_agreement_studio.import_engine.workbook import inspect_workbook
from metaphor_agreement_studio.persistence.project_service import create_project, record_analysis_version, save_dataset_version
from metaphor_agreement_studio.statistics.engine import analyze_dataset
from metaphor_agreement_studio.statistics.types import AnalysisConfig
from metaphor_agreement_studio.validation.rater_growth import (
    RaterMergeBlocked,
    align_validated_datasets,
    merge_additional_raters,
)
from metaphor_agreement_studio.validation.service import ValidationService
from tests.fixtures.workbook_factory import build_two_rater_workbook


def _decisions(draft):
    output = []
    for index, issue in enumerate(draft.issues, start=1):
        if not issue.blocks_validation:
            continue
        if issue.code == "category_normalization_decision":
            value = "Adjective"
        elif issue.code in {"confirm_annotation_mapping", "annotation_mapping_conflict"}:
            value = "confirm_green_metaphor_red_non_metaphor"
        elif issue.code == "aggregate_role_decision":
            value = "aggregate"
        else:
            raise AssertionError(issue.code)
        output.append(
            ValidationDecision(
                decision_id=f"d{index}",
                decision_type=issue.decision_type or "",
                target_id=issue.target_id or "",
                original_value=None,
                validated_value=value,
                reason="Phase 7 rater growth test",
                decided_at="2026-09-06T00:05:00+01:00",
            )
        )
    return tuple(output)


def _validate(path: Path):
    service = ValidationService()
    draft = service.prepare(inspect_workbook(path))
    return service.validate(draft, _decisions(draft))


def _add_third_rater_column(source: Path, destination: Path) -> Path:
    wb = load_workbook(source)
    ws = wb["Planilha1"]
    ws["N26"] = "Sofia"
    ws["N27"] = "Is metaphor?"
    green = PatternFill("solid", fgColor="92D050")
    red = PatternFill("solid", fgColor="F4CCCC")
    for coordinate, label, fill in (
        ("N28", "YES", green),
        ("N29", "NO", red),
        ("N30", "YES", green),
    ):
        ws[coordinate] = label
        ws[coordinate].fill = fill
    wb.save(destination)
    wb.close()
    return destination


def _build_separate_sofia_file(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws["I24"] = "1. Video: Example"
    ws["L25"] = "Raters"
    ws["L26"] = "Sofia"
    for idx, value in enumerate(
        ["Nr", "Lexical unit", "Grammatical category", "Is metaphor?"], start=9
    ):
        ws.cell(27, idx, value)
    green = PatternFill("solid", fgColor="92D050")
    red = PatternFill("solid", fgColor="F4CCCC")
    rows = [
        (1, "top", "Adjetive", "YES", green),
        (2, "hand!", "Noun", "NO", red),
        (3, "spider", "Noun", "YES", green),
    ]
    for row_index, (order, lexical, pos, label, fill) in enumerate(rows, start=28):
        ws.cell(row_index, 9, order)
        ws.cell(row_index, 10, lexical)
        ws.cell(row_index, 11, pos)
        ws.cell(row_index, 12, label)
        ws.cell(row_index, 12).fill = fill
    wb.save(path)
    wb.close()
    return path


def test_same_workbook_third_rater_creates_new_dataset_version_and_multirater_results(
    tmp_path: Path,
) -> None:
    two_path = build_two_rater_workbook(tmp_path / "two.xlsx")
    two = _validate(two_path)
    project = create_project(tmp_path / "project", "Rater Growth", two, (two_path,))
    project = record_analysis_version(project, AnalysisConfig(bootstrap_samples=0))
    old_analysis_id = project.current_analysis_version.analysis_id

    three_path = _add_third_rater_column(two_path, tmp_path / "three.xlsx")
    three = _validate(three_path)
    project = save_dataset_version(project, three, "Third rater added", (three_path,))

    assert len(three.raters) == 3
    assert project.current_dataset_version.ordinal == 2
    assert project.current_analysis_version.analysis_id == old_analysis_id
    assert project.current_analysis_version.dataset_version_id != project.current_dataset_version.version_id

    bundle = analyze_dataset(three, AnalysisConfig(bootstrap_samples=0))
    assert len(bundle.pairwise) == 3
    assert {metric.metric for metric in bundle.overall_multirater} >= {"fleiss_kappa", "krippendorff_alpha"}


def test_separate_rater_file_requires_probable_match_confirmation_before_merge(
    tmp_path: Path,
) -> None:
    reference_path = build_two_rater_workbook(tmp_path / "reference.xlsx")
    reference = _validate(reference_path)
    incoming_path = _build_separate_sofia_file(tmp_path / "sofia.xlsx")
    incoming = _validate(incoming_path)

    alignment = align_validated_datasets(reference, incoming)
    assert len(alignment.exact) == 2
    assert len(alignment.probable) == 1
    probable_id = alignment.probable[0].incoming.unit_id

    with pytest.raises(RaterMergeBlocked):
        merge_additional_raters(reference, incoming, alignment)

    merged = merge_additional_raters(
        reference,
        incoming,
        alignment,
        accepted_probable_incoming_ids=(probable_id,),
    )
    assert {r.display_name for r in merged.raters} == {"Eduardo", "Braulio", "Sofia"}
    bundle = analyze_dataset(merged, AnalysisConfig(bootstrap_samples=0))
    assert len(bundle.pairwise) == 3

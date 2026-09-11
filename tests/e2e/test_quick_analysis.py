from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from metaphor_agreement_studio.domain.imports import ValidationDecision
from metaphor_agreement_studio.export.workbook import REQUIRED_SHEETS, export_results_workbook
from metaphor_agreement_studio.import_engine.discovery import discover_workbooks
from metaphor_agreement_studio.import_engine.workbook import inspect_workbook
from metaphor_agreement_studio.review.service import build_review_cases
from metaphor_agreement_studio.review.types import ReviewStatus
from metaphor_agreement_studio.statistics.engine import analyze_dataset
from metaphor_agreement_studio.statistics.types import AnalysisConfig
from metaphor_agreement_studio.validation.service import ValidationBlocked, ValidationService
from tests.fixtures.workbook_factory import build_tables_with_aggregate


def _build_quick_analysis_workbook(path: Path) -> Path:
    build_tables_with_aggregate(path)
    wb = load_workbook(path)
    ws = wb["Ratings"]
    green = PatternFill("solid", fgColor="92D050")
    red = PatternFill("solid", fgColor="F4CCCC")
    blank = PatternFill(fill_type=None)

    # Primary source A: one disagreement and one unanimous non-metaphor.
    ws["E6"] = "NO"
    ws["E6"].fill = red
    ws["F6"] = "YES"
    ws["F6"].fill = green
    ws["E7"] = "NO"
    ws["E7"].fill = red
    ws["F7"] = "NO"
    ws["F7"].fill = red

    # Primary source B: typo proposal + missing rating.
    ws["K14"] = "Adjetive"
    ws["L14"] = "YES"
    ws["L14"].fill = green
    ws["M14"] = None
    ws["M14"].fill = blank

    # Mirror the aggregate source so it remains a complete aggregate candidate.
    ws["R24"] = "Adjetive"
    ws["S22"] = "NO"
    ws["S22"].fill = red
    ws["T22"] = "YES"
    ws["T22"].fill = green
    ws["S23"] = "NO"
    ws["S23"].fill = red
    ws["T23"] = "NO"
    ws["T23"].fill = red
    ws["S24"] = "YES"
    ws["S24"].fill = green
    ws["T24"] = None
    ws["T24"].fill = blank

    wb.save(path)
    wb.close()
    return path


def _researcher_decisions(draft) -> tuple[ValidationDecision, ...]:
    decisions: list[ValidationDecision] = []
    for index, issue in enumerate(draft.issues, start=1):
        if not issue.blocks_validation:
            continue
        if issue.code == "category_normalization_decision":
            value = "Adjective"
        elif issue.code in {"confirm_annotation_mapping", "annotation_mapping_conflict"}:
            value = "confirm_green_metaphor_red_non_metaphor"
        elif issue.code == "aggregate_role_decision":
            value = "aggregate"
        else:
            raise AssertionError(f"Unhandled blocking issue: {issue.code}")
        decisions.append(
            ValidationDecision(
                decision_id=f"phase7-{index}",
                decision_type=issue.decision_type or "",
                target_id=issue.target_id or "",
                original_value=None,
                validated_value=value,
                reason="Phase 7 end-to-end acceptance decision",
                decided_at="2026-09-05T23:45:00+01:00",
            )
        )
    return tuple(decisions)


def test_complete_quick_analysis_workflow(tmp_path: Path) -> None:
    workbook_path = _build_quick_analysis_workbook(tmp_path / "quick-analysis.xlsx")

    candidates = discover_workbooks((tmp_path,))
    assert [item.display_name for item in candidates] == ["quick-analysis.xlsx"]

    service = ValidationService()
    draft = service.prepare(inspect_workbook(workbook_path))
    assert len(draft.tables) == 3
    assert {item.display_name for item in draft.rater_mappings} == {"Eduardo", "Braulio"}
    assert len(draft.aggregate_candidates) == 1
    assert any(item.original == "Adjetive" and item.proposed == "Adjective" for item in draft.normalization_proposals)

    with pytest.raises(ValidationBlocked):
        service.validate(draft, ())

    dataset = service.validate(draft, _researcher_decisions(draft))
    analytical_sources = {source.source_id for source in dataset.sources if not source.is_aggregate}
    analytical_units = [unit for unit in dataset.units if unit.source_id in analytical_sources]
    assert len(analytical_units) == 3

    bundle = analyze_dataset(dataset, AnalysisConfig(bootstrap_samples=0))
    assert len(bundle.pairwise) == 1
    assert bundle.cochran_q is not None

    review = build_review_cases(
        dataset,
        unit_ids=tuple(unit.unit_id for unit in analytical_units),
    )
    statuses = {case.status for case in review}
    assert ReviewStatus.DISAGREEMENT in statuses
    assert ReviewStatus.UNANIMOUS_NON_METAPHOR in statuses
    assert any(case.has_missing for case in review)

    output = export_results_workbook(dataset, bundle, (), tmp_path / "results.xlsx")
    assert output.exists()
    exported = load_workbook(output, read_only=True)
    try:
        assert tuple(exported.sheetnames) == REQUIRED_SHEETS
    finally:
        exported.close()

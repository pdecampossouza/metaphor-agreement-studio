from __future__ import annotations

import hashlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from metaphor_agreement_studio.domain.imports import ValidationDecision
from metaphor_agreement_studio.import_engine.workbook import inspect_workbook
from metaphor_agreement_studio.persistence.project_service import create_project, save_dataset_version
from metaphor_agreement_studio.validation.rater_growth import align_validated_datasets, merge_additional_raters
from metaphor_agreement_studio.validation.service import ValidationService
from tests.fixtures.workbook_factory import build_two_rater_workbook


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _decisions(draft) -> tuple[ValidationDecision, ...]:
    decisions: list[ValidationDecision] = []
    for index, issue in enumerate(draft.issues, start=1):
        if not issue.blocks_validation:
            continue
        if issue.code == "category_normalization_decision":
            value = "Adjective"
        elif issue.code in {"confirm_annotation_mapping", "annotation_mapping_conflict"}:
            value = "confirm_green_metaphor_red_non_metaphor"
        elif issue.code == "aggregate_role_decision":
            value = "aggregate"
        else:
            raise AssertionError(f"Unhandled blocking issue: {issue.code}")
        decisions.append(
            ValidationDecision(
                decision_id=f"immutability-{index}",
                decision_type=issue.decision_type or "",
                target_id=issue.target_id or "",
                original_value=None,
                validated_value=value,
                reason="Release source immutability gate",
                decided_at="2026-09-06T00:30:00+01:00",
            )
        )
    return tuple(decisions)


def _validate(path: Path):
    service = ValidationService()
    draft = service.prepare(inspect_workbook(path))
    return service.validate(draft, _decisions(draft))


def _build_separate_rater_file(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Planilha1"
    sheet["I24"] = "1. Video: Example"
    sheet["L25"] = "Raters"
    sheet["L26"] = "Sofia"
    for column, value in enumerate(
        ["Nr", "Lexical unit", "Grammatical category", "Is metaphor?"], start=9
    ):
        sheet.cell(27, column, value)
    green = PatternFill("solid", fgColor="92D050")
    red = PatternFill("solid", fgColor="F4CCCC")
    for row, values in enumerate(
        (
            (1, "top", "Adjetive", "YES", green),
            (2, "hand", "Noun", "NO", red),
            (3, "spider", "Noun", "YES", green),
        ),
        start=28,
    ):
        order, lexical, category, label, fill = values
        sheet.cell(row, 9, order)
        sheet.cell(row, 10, lexical)
        sheet.cell(row, 11, category)
        sheet.cell(row, 12, label)
        sheet.cell(row, 12).fill = fill
    workbook.save(path)
    workbook.close()
    return path


def test_source_workbooks_remain_byte_identical_across_all_ingest_service_paths(
    tmp_path: Path,
) -> None:
    primary_path = build_two_rater_workbook(tmp_path / "primary.xlsx")
    incoming_path = _build_separate_rater_file(tmp_path / "sofia.xlsx")
    before = {primary_path: _sha256(primary_path), incoming_path: _sha256(incoming_path)}

    # Inspect/validate the original source without writing to it.
    primary = _validate(primary_path)

    # Persisting a Research Project must copy the original into project storage, never mutate it.
    project = create_project(tmp_path / "project", "Immutability Gate", primary, (primary_path,))

    # A separately supplied rater workbook follows the same inspect/validate/alignment path.
    incoming = _validate(incoming_path)
    alignment = align_validated_datasets(primary, incoming)
    merged = merge_additional_raters(primary, incoming, alignment)
    save_dataset_version(project, merged, "Validated rater added", (incoming_path,))

    after = {primary_path: _sha256(primary_path), incoming_path: _sha256(incoming_path)}
    assert after == before

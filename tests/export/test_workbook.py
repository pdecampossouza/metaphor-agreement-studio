from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from metaphor_agreement_studio.domain.enums import Classification, ValidationStatus
from metaphor_agreement_studio.domain.imports import ValidatedDataset
from metaphor_agreement_studio.domain.models import Annotation, LexicalUnit, Rater, Source
from metaphor_agreement_studio.export.datasets import export_long_csv
from metaphor_agreement_studio.export.workbook import REQUIRED_SHEETS, export_results_workbook
from metaphor_agreement_studio.statistics.engine import analyze_dataset
from metaphor_agreement_studio.statistics.types import AnalysisConfig


def dataset() -> ValidatedDataset:
    units = (
        LexicalUnit("u1", "spider", "Noun", "s1", 1),
        LexicalUnit("u2", "hand", "Noun", "s1", 1),
        LexicalUnit("u3", "move", "Verb", "s1", 1),
    )
    raters = (Rater("r1", "Eduardo"), Rater("r2", "Braulio"))
    values = {
        ("u1", "r1"): Classification.METAPHOR,
        ("u1", "r2"): Classification.METAPHOR,
        ("u2", "r1"): Classification.NON_METAPHOR,
        ("u2", "r2"): Classification.NON_METAPHOR,
        ("u3", "r1"): Classification.METAPHOR,
        ("u3", "r2"): Classification.NON_METAPHOR,
    }
    annotations = tuple(
        Annotation(
            f"a{i}", unit_id, rater_id, value, "imp", "Sheet1", f"D{i}", value.value, {}, "fixture", ValidationStatus.VALIDATED
        )
        for i, ((unit_id, rater_id), value) in enumerate(values.items(), start=1)
    )
    return ValidatedDataset(
        "d1", (Source("s1", "Example"),), units, raters, annotations, (), ()
    )


def test_results_workbook_has_required_sheets_in_exact_order(tmp_path: Path) -> None:
    ds = dataset()
    bundle = analyze_dataset(ds, AnalysisConfig(selected_rater_ids=("r1", "r2"), bootstrap_samples=20))
    out = export_results_workbook(ds, bundle, (), tmp_path / "Metaphor_Agreement_Results.xlsx")
    wb = load_workbook(out, data_only=True)
    assert wb.sheetnames == list(REQUIRED_SHEETS)


def test_readme_documents_conventions_and_context(tmp_path: Path) -> None:
    ds = dataset()
    bundle = analyze_dataset(ds, AnalysisConfig(selected_rater_ids=("r1", "r2"), bootstrap_samples=20))
    out = export_results_workbook(
        ds,
        bundle,
        (),
        tmp_path / "results.xlsx",
        project_name="Metaphor Study",
        dataset_version="v1.0",
    )
    wb = load_workbook(out, data_only=True)
    values = [cell.value for row in wb["README"].iter_rows() for cell in row if cell.value is not None]
    joined = "\n".join(str(v) for v in values)
    assert "Metaphor Study" in joined
    assert "v1.0" in joined
    assert "Eduardo" in joined and "Braulio" in joined
    assert "Metaphor = 1" in joined
    assert "Non-metaphor = 0" in joined
    assert "Missing = NA" in joined


def test_machine_values_are_numeric_or_documented_text_status(tmp_path: Path) -> None:
    ds = dataset()
    bundle = analyze_dataset(ds, AnalysisConfig(selected_rater_ids=("r1", "r2"), bootstrap_samples=20))
    out = export_results_workbook(ds, bundle, (), tmp_path / "results.xlsx")
    wb = load_workbook(out, data_only=False)
    ws = wb["Pairwise Kappa"]
    headers = {cell.value: cell.column for cell in ws[1]}
    kappa = ws.cell(2, headers["Cohen's Kappa"]).value
    n = ws.cell(2, headers["Effective N"]).value
    assert isinstance(kappa, (int, float)) or isinstance(kappa, str)
    assert isinstance(n, int)
    assert not (isinstance(kappa, str) and kappa.startswith("="))


def test_long_csv_uses_binary_machine_conventions(tmp_path: Path) -> None:
    out = export_long_csv(dataset(), tmp_path / "annotations_long.csv")
    text = out.read_text(encoding="utf-8")
    assert "classification_binary" in text
    assert ",1," in text or text.rstrip().endswith(",1")
    assert ",0," in text or text.rstrip().endswith(",0")


def _multirater_dataset() -> ValidatedDataset:
    base = dataset()
    third = Rater("r3", "Sofia")
    third_values = {
        "u1": Classification.METAPHOR,
        "u2": Classification.METAPHOR,
        "u3": Classification.NON_METAPHOR,
    }
    extra = tuple(
        Annotation(
            f"m{i}", unit_id, "r3", value, "imp", "Sheet1", f"M{i}", value.value, {}, "fixture", ValidationStatus.VALIDATED
        )
        for i, (unit_id, value) in enumerate(third_values.items(), start=1)
    )
    return ValidatedDataset(
        "d3",
        base.sources,
        base.units,
        base.raters + (third,),
        base.annotations + extra,
        (),
        (),
    )


def test_group_export_includes_fleiss_kappa_for_three_or_more_raters(tmp_path: Path) -> None:
    ds = _multirater_dataset()
    bundle = analyze_dataset(ds, AnalysisConfig(bootstrap_samples=0))
    out = export_results_workbook(ds, bundle, (), tmp_path / "results-multirater.xlsx")
    wb = load_workbook(out, data_only=True)

    for sheet_name in ("By POS", "By Source"):
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        assert "Fleiss' Kappa" in headers
        fleiss_col = headers.index("Fleiss' Kappa") + 1
        assert any(ws.cell(row, fleiss_col).value is not None for row in range(2, ws.max_row + 1))

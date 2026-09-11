import hashlib
from pathlib import Path

from metaphor_agreement_studio.import_engine.workbook import inspect_workbook
from tests.fixtures.workbook_factory import build_two_rater_workbook


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_inspection_preserves_value_coordinate_and_fill(tmp_path: Path) -> None:
    path = build_two_rater_workbook(tmp_path / "ratings.xlsx")
    before = sha256(path)

    inspection = inspect_workbook(path)

    assert inspection.path == path.resolve()
    assert inspection.sheet_names == ("Planilha1",)
    cell = inspection.sheet("Planilha1").cell("L30")
    assert cell.raw_value == "YES"
    assert cell.coordinate == "L30"
    assert cell.fill.fill_type == "solid"
    assert cell.fill.fg_rgb in {"00C6E0B4", "C6E0B4", "FFC6E0B4"}
    assert sha256(path) == before


def test_inspection_preserves_merged_ranges(tmp_path: Path) -> None:
    path = build_two_rater_workbook(tmp_path / "ratings.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["Planilha1"]
    ws.merge_cells("I24:K24")
    wb.save(path)

    inspection = inspect_workbook(path)

    assert "I24:K24" in inspection.sheet("Planilha1").merged_ranges

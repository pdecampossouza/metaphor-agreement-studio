from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

GREEN_A = "92D050"
GREEN_B = "C6E0B4"
RED = "F4CCCC"


def build_two_rater_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws["I24"] = "1. Video: Example"
    ws["L25"] = "Raters"
    ws["L26"] = "Eduardo"
    ws["M26"] = "Braulio"
    headers = ["Nr", "Lexical unit", "Grammatical category", "Is metaphor?", "Is metaphor?"]
    for idx, value in enumerate(headers, start=9):
        ws.cell(27, idx, value)
    rows = [
        (1, "top", "Adjetive", None, None),
        (2, "hand", "Noun", "NO", None),
        (3, "spider", "Noun", "YES", None),
    ]
    for r, row in enumerate(rows, start=28):
        for c, value in enumerate(row, start=9):
            ws.cell(r, c, value)
    ws["L28"].fill = PatternFill("solid", fgColor=RED)
    ws["M28"].fill = PatternFill("solid", fgColor=GREEN_A)
    ws["L29"].fill = PatternFill("solid", fgColor=RED)
    ws["M29"].fill = PatternFill("solid", fgColor=RED)
    ws["L30"].fill = PatternFill("solid", fgColor=GREEN_B)
    ws["M30"].fill = PatternFill("solid", fgColor=GREEN_A)
    wb.save(path)
    return path


def build_conflicting_mapping_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws["A1"] = "1. Video: Conflict"
    ws["D2"] = "Raters"
    ws["D3"] = "Eduardo"
    headers = ["Nr", "Lexical unit", "Grammatical category", "Is metaphor?"]
    for idx, value in enumerate(headers, start=1):
        ws.cell(4, idx, value)
    ws["A5"] = 1
    ws["B5"] = "alpha"
    ws["C5"] = "Noun"
    ws["D5"] = "YES"
    ws["D5"].fill = PatternFill("solid", fgColor=RED)
    ws["A6"] = 2
    ws["B6"] = "beta"
    ws["C6"] = "Noun"
    ws["D6"] = "NO"
    ws["D6"].fill = PatternFill("solid", fgColor=GREEN_A)
    wb.save(path)
    return path


def build_tables_with_aggregate(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ratings"

    def add_table(start_row: int, start_col: int, title: str, rows: list[tuple[int, str, str]]) -> None:
        ws.cell(start_row, start_col, title)
        ws.cell(start_row + 1, start_col + 3, "Raters")
        ws.cell(start_row + 2, start_col + 3, "Eduardo")
        ws.cell(start_row + 2, start_col + 4, "Braulio")
        headers = ["Nr", "Lexical unit", "Grammatical category", "Is metaphor?", "Is metaphor?"]
        for offset, value in enumerate(headers):
            ws.cell(start_row + 3, start_col + offset, value)
        for row_offset, (number, lexical, pos) in enumerate(rows, start=4):
            row = start_row + row_offset
            ws.cell(row, start_col, number)
            ws.cell(row, start_col + 1, lexical)
            ws.cell(row, start_col + 2, pos)
            ws.cell(row, start_col + 3).fill = PatternFill("solid", fgColor=RED)
            ws.cell(row, start_col + 4).fill = PatternFill("solid", fgColor=GREEN_A)

    a = [(1, "alpha", "Noun"), (2, "beta", "Verb")]
    b = [(1, "gamma", "Adjective")]
    combined = [(1, "alpha", "Noun"), (2, "beta", "Verb"), (3, "gamma", "Adjective")]
    add_table(2, 2, "1. Video: A", a)
    add_table(10, 9, "2. Video: B", b)
    add_table(18, 16, "3. Combined Videos", combined)
    wb.save(path)
    return path

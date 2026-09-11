from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from metaphor_agreement_studio.domain.enums import Classification
from metaphor_agreement_studio.domain.imports import ValidatedDataset
from metaphor_agreement_studio.statistics.types import AnalysisBundle

REQUIRED_SHEETS = (
    "README",
    "Dataset",
    "Annotations",
    "Rater Summary",
    "Overall Statistics",
    "Pairwise Kappa",
    "Cochran Q",
    "By POS",
    "By Source",
    "Disagreements",
    "Unanimous Cases",
    "Missing Data",
    "Data Quality",
    "Audit Log",
)

_HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
_HEADER_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)


def _status_value(metric) -> float | str:
    return metric.value if metric.value is not None else metric.status.value.replace("_", " ")


def _write_table(ws, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> None:
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append(list(row))
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for column in range(1, ws.max_column + 1):
        width = max(
            10,
            min(
                42,
                max(
                    (len(str(ws.cell(row, column).value or "")) for row in range(1, ws.max_row + 1)),
                    default=10,
                )
                + 2,
            ),
        )
        ws.column_dimensions[get_column_letter(column)].width = width


def _unit_rows(dataset: ValidatedDataset):
    source_map = {s.source_id: s for s in dataset.sources}
    for unit in dataset.units:
        source = source_map[unit.source_id]
        yield (
            unit.unit_id,
            unit.lexical_unit,
            unit.grammatical_category,
            source.display_name,
            source.is_aggregate,
            unit.occurrence_index,
            unit.song,
            unit.artist,
            unit.verse,
            unit.context,
            unit.timestamp,
        )


def _annotation_rows(dataset: ValidatedDataset):
    unit_map = {u.unit_id: u for u in dataset.units}
    source_map = {s.source_id: s for s in dataset.sources}
    rater_map = {r.rater_id: r for r in dataset.raters}
    for ann in dataset.annotations:
        unit = unit_map[ann.unit_id]
        binary = 1 if ann.classification is Classification.METAPHOR else 0 if ann.classification is Classification.NON_METAPHOR else "NA"
        yield (
            ann.annotation_id,
            unit.lexical_unit,
            unit.grammatical_category,
            source_map[unit.source_id].display_name,
            rater_map[ann.rater_id].display_name,
            ann.classification.value,
            binary,
            ann.original_sheet,
            ann.original_cell,
            ann.detection_method,
        )


def _rater_summary_rows(dataset: ValidatedDataset):
    non_aggregate = {s.source_id for s in dataset.sources if not s.is_aggregate}
    unit_ids = {u.unit_id for u in dataset.units if u.source_id in non_aggregate}
    for rater in dataset.raters:
        values = [a.classification for a in dataset.annotations if a.rater_id == rater.rater_id and a.unit_id in unit_ids]
        metaphor = sum(v is Classification.METAPHOR for v in values)
        non_metaphor = sum(v is Classification.NON_METAPHOR for v in values)
        missing = sum(v is Classification.MISSING for v in values)
        observed = metaphor + non_metaphor
        yield (rater.display_name, len(unit_ids), metaphor, non_metaphor, missing, metaphor / observed if observed else None)


def _group_rows(groups):
    for group in groups:
        fleiss = _status_value(group.fleiss_kappa) if group.fleiss_kappa is not None else "NA"
        if not group.pairwise:
            yield (group.display_name, group.lexical_unit_count, "", "", "", fleiss, "", "", "")
        for pair in group.pairwise:
            q = group.cochran_q
            yield (
                group.display_name,
                group.lexical_unit_count,
                f"{pair.rater_a_id} × {pair.rater_b_id}",
                pair.raw_agreement.value if pair.raw_agreement.value is not None else pair.raw_agreement.status.value,
                _status_value(pair.kappa),
                fleiss,
                q.q if q and q.q is not None else q.status.value if q else "NA",
                q.degrees_of_freedom if q else None,
                q.p_value if q else None,
            )


def _case_rows(dataset: ValidatedDataset, *, disagreement: bool):
    non_aggregate = {s.source_id for s in dataset.sources if not s.is_aggregate}
    source_map = {s.source_id: s.display_name for s in dataset.sources}
    raters = tuple(dataset.raters)
    ann = {(a.unit_id, a.rater_id): a.classification for a in dataset.annotations}
    for unit in dataset.units:
        if unit.source_id not in non_aggregate:
            continue
        values = [ann.get((unit.unit_id, r.rater_id), Classification.MISSING) for r in raters]
        observed = [value for value in values if value is not Classification.MISSING]
        is_disagreement = len(set(observed)) > 1
        is_unanimous = len(observed) >= 2 and len(set(observed)) == 1
        if disagreement and not is_disagreement:
            continue
        if not disagreement and not is_unanimous:
            continue
        yield (
            unit.lexical_unit,
            unit.grammatical_category,
            source_map[unit.source_id],
            *[value.value for value in values],
        )


def export_results_workbook(
    dataset: ValidatedDataset,
    analysis_bundle: AnalysisBundle,
    audit_events,
    destination: Path,
    *,
    project_name: str = "Quick Analysis",
    dataset_version: str = "Quick Analysis",
    analysis_version: str = "Current session",
) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name in REQUIRED_SHEETS:
        wb.create_sheet(name)

    readme = wb["README"]
    readme["A1"] = "METAPHOR AGREEMENT STUDIO"
    readme["A1"].font = _TITLE_FONT
    items = [
        ("Project / mode", project_name),
        ("Dataset version", dataset_version),
        ("Analysis version", analysis_version),
        ("Generated (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Raters", ", ".join(r.display_name for r in dataset.raters)),
        ("Lexical units", analysis_bundle.counts.lexical_units),
        ("Convention", "Metaphor = 1"),
        ("Convention", "Non-metaphor = 0"),
        ("Convention", "Missing = NA"),
    ]
    for row, (label, value) in enumerate(items, start=3):
        readme.cell(row, 1, label).font = _HEADER_FONT
        readme.cell(row, 2, value)
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 70

    _write_table(
        wb["Dataset"],
        ["Unit ID", "Lexical unit", "Grammatical category", "Source", "Aggregate", "Occurrence", "Song", "Artist", "Verse", "Context", "Timestamp"],
        _unit_rows(dataset),
    )
    _write_table(
        wb["Annotations"],
        ["Annotation ID", "Lexical unit", "Grammatical category", "Source", "Rater", "Classification", "Binary", "Original sheet", "Original cell", "Detection method"],
        _annotation_rows(dataset),
    )
    _write_table(
        wb["Rater Summary"],
        ["Rater", "Analytical units", "Metaphor", "Non-metaphor", "Missing", "Metaphor rate"],
        _rater_summary_rows(dataset),
    )
    q = analysis_bundle.cochran_q
    overall_rows = [
        ("Lexical units", analysis_bundle.counts.lexical_units, "count"),
        ("Raters", analysis_bundle.counts.raters, "count"),
        ("Metaphor ratings", analysis_bundle.counts.metaphor_ratings, "count"),
        ("Non-metaphor ratings", analysis_bundle.counts.non_metaphor_ratings, "count"),
        ("Missing ratings", analysis_bundle.counts.missing_ratings, "count"),
    ]
    for metric in analysis_bundle.overall_multirater:
        overall_rows.append((metric.metric, _status_value(metric), metric.status.value))
    _write_table(wb["Overall Statistics"], ["Metric", "Value", "Status"], overall_rows)

    _write_table(
        wb["Pairwise Kappa"],
        ["Rater A", "Rater B", "Effective N", "Raw agreement", "Cohen's Kappa", "CI low", "CI high", "Status"],
        [
            (
                p.rater_a_id,
                p.rater_b_id,
                p.kappa.effective_n,
                p.raw_agreement.value if p.raw_agreement.value is not None else p.raw_agreement.status.value,
                _status_value(p.kappa),
                p.kappa.ci_low,
                p.kappa.ci_high,
                p.kappa.status.value,
            )
            for p in analysis_bundle.pairwise
        ],
    )
    _write_table(
        wb["Cochran Q"],
        ["Scope", "Q", "df", "p-value", "Effective N", "Status"],
        [
            (
                "Overall",
                q.q if q and q.q is not None else q.status.value if q else "NA",
                q.degrees_of_freedom if q else None,
                q.p_value if q else None,
                q.effective_n if q else 0,
                q.status.value if q else "not_available",
            )
        ],
    )
    headers = ["Group", "Lexical units", "Rater pair", "Raw agreement", "Cohen's Kappa", "Fleiss' Kappa", "Cochran Q", "Q df", "Q p-value"]
    _write_table(wb["By POS"], headers, _group_rows(analysis_bundle.by_category))
    _write_table(wb["By Source"], headers, _group_rows(analysis_bundle.by_source))

    case_headers = ["Lexical unit", "Grammatical category", "Source", *[r.display_name for r in dataset.raters]]
    _write_table(wb["Disagreements"], case_headers, _case_rows(dataset, disagreement=True))
    _write_table(wb["Unanimous Cases"], case_headers, _case_rows(dataset, disagreement=False))

    missing_rows = []
    for ann in dataset.annotations:
        if ann.classification is Classification.MISSING:
            missing_rows.append((ann.unit_id, ann.rater_id, ann.original_sheet, ann.original_cell))
    _write_table(wb["Missing Data"], ["Unit ID", "Rater ID", "Original sheet", "Original cell"], missing_rows)

    quality_rows = []
    for note in dataset.quality_notes:
        quality_rows.append((getattr(note, "code", "quality_note"), getattr(note, "severity", ""), getattr(note, "message", str(note))))
    for decision in dataset.validation_decisions:
        quality_rows.append((decision.decision_type, "validated", f"{decision.original_value} → {decision.validated_value}: {decision.reason}"))
    _write_table(wb["Data Quality"], ["Type", "Status", "Description"], quality_rows)

    audit_rows = [
        (
            getattr(event, "occurred_at", ""),
            getattr(event, "event_type", ""),
            getattr(event, "summary", str(event)),
            getattr(event, "dataset_version_id", None),
            getattr(event, "analysis_id", None),
        )
        for event in audit_events
    ]
    _write_table(wb["Audit Log"], ["Occurred at", "Event type", "Summary", "Dataset version", "Analysis version"], audit_rows)

    wb.save(destination)
    return destination

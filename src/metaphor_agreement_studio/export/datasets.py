from __future__ import annotations

import csv
from pathlib import Path

from metaphor_agreement_studio.domain.enums import Classification
from metaphor_agreement_studio.domain.imports import ValidatedDataset


def _binary(classification: Classification):
    if classification is Classification.METAPHOR:
        return 1
    if classification is Classification.NON_METAPHOR:
        return 0
    return "NA"


def export_long_csv(dataset: ValidatedDataset, destination: Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    units = {unit.unit_id: unit for unit in dataset.units}
    sources = {source.source_id: source for source in dataset.sources}
    raters = {rater.rater_id: rater for rater in dataset.raters}
    fields = [
        "unit_id",
        "lexical_unit",
        "grammatical_category",
        "source_id",
        "source",
        "is_aggregate",
        "occurrence_index",
        "rater_id",
        "rater",
        "classification",
        "classification_binary",
        "original_sheet",
        "original_cell",
        "detection_method",
    ]
    with destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for annotation in dataset.annotations:
            unit = units[annotation.unit_id]
            source = sources[unit.source_id]
            rater = raters[annotation.rater_id]
            writer.writerow(
                {
                    "unit_id": unit.unit_id,
                    "lexical_unit": unit.lexical_unit,
                    "grammatical_category": unit.grammatical_category,
                    "source_id": source.source_id,
                    "source": source.display_name,
                    "is_aggregate": source.is_aggregate,
                    "occurrence_index": unit.occurrence_index,
                    "rater_id": rater.rater_id,
                    "rater": rater.display_name,
                    "classification": annotation.classification.value,
                    "classification_binary": _binary(annotation.classification),
                    "original_sheet": annotation.original_sheet,
                    "original_cell": annotation.original_cell,
                    "detection_method": annotation.detection_method,
                }
            )
    return destination


def export_validated_xlsx(dataset: ValidatedDataset, destination: Path) -> Path:
    """Export validated analytical data without modifying any source workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Dataset"
    headers = ["unit_id", "lexical_unit", "grammatical_category", "source_id", "occurrence_index", "song", "artist", "verse", "context", "timestamp"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
    for unit in dataset.units:
        ws.append([
            unit.unit_id,
            unit.lexical_unit,
            unit.grammatical_category,
            unit.source_id,
            unit.occurrence_index,
            unit.song,
            unit.artist,
            unit.verse,
            unit.context,
            unit.timestamp,
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    anns = wb.create_sheet("Annotations")
    ann_headers = ["annotation_id", "unit_id", "rater_id", "classification", "classification_binary", "original_sheet", "original_cell", "detection_method"]
    anns.append(ann_headers)
    for cell in anns[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
    for ann in dataset.annotations:
        anns.append([
            ann.annotation_id,
            ann.unit_id,
            ann.rater_id,
            ann.classification.value,
            _binary(ann.classification),
            ann.original_sheet,
            ann.original_cell,
            ann.detection_method,
        ])
    anns.freeze_panes = "A2"
    anns.auto_filter.ref = anns.dimensions
    wb.save(destination)
    return destination

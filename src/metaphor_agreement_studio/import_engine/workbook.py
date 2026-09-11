from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from metaphor_agreement_studio.domain.imports import (
    CellSnapshot,
    FillSnapshot,
    SheetSnapshot,
    WorkbookInspection,
)


def _color_value(color: Any, name: str, expected_type: type) -> Any:
    value = getattr(color, name, None)
    return value if isinstance(value, expected_type) else None


def _fill_snapshot(cell: Any) -> FillSnapshot:
    fill = getattr(cell, "fill", None)
    fg = getattr(fill, "fgColor", None)
    bg = getattr(fill, "bgColor", None)
    return FillSnapshot(
        fill_type=getattr(fill, "fill_type", None),
        fg_type=getattr(fg, "type", None),
        fg_rgb=_color_value(fg, "rgb", str),
        fg_indexed=_color_value(fg, "indexed", int),
        fg_theme=_color_value(fg, "theme", int),
        fg_tint=float(getattr(fg, "tint", 0.0) or 0.0),
        bg_type=getattr(bg, "type", None),
        bg_rgb=_color_value(bg, "rgb", str),
        bg_indexed=_color_value(bg, "indexed", int),
        bg_theme=_color_value(bg, "theme", int),
        bg_tint=float(getattr(bg, "tint", 0.0) or 0.0),
    )


def _snapshot_cell(path: Path, sheet_name: str, cell: Any) -> CellSnapshot:
    style = getattr(cell, "_style", None)
    return CellSnapshot(
        workbook_path=path,
        sheet_name=sheet_name,
        coordinate=str(getattr(cell, "coordinate")),
        row=int(getattr(cell, "row")),
        column=int(getattr(cell, "column")),
        raw_value=getattr(cell, "value", None),
        fill=_fill_snapshot(cell),
        style_id=getattr(cell, "style_id", None),
        font_id=getattr(style, "fontId", None),
        border_id=getattr(style, "borderId", None),
        number_format=getattr(cell, "number_format", None),
    )


def inspect_workbook(path: Path) -> WorkbookInspection:
    resolved = Path(path).expanduser().resolve(strict=True)
    keep_vba = resolved.suffix.casefold() == ".xlsm"
    workbook = load_workbook(
        resolved,
        data_only=False,
        read_only=False,
        keep_vba=keep_vba,
    )
    try:
        sheets: list[SheetSnapshot] = []
        for worksheet in workbook.worksheets:
            cells = tuple(
                _snapshot_cell(resolved, worksheet.title, cell)
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=worksheet.max_row,
                    min_col=1,
                    max_col=worksheet.max_column,
                )
                for cell in row
            )
            sheets.append(
                SheetSnapshot(
                    name=worksheet.title,
                    cells=cells,
                    max_row=worksheet.max_row,
                    max_column=worksheet.max_column,
                    merged_ranges=tuple(str(item) for item in worksheet.merged_cells.ranges),
                )
            )
        return WorkbookInspection(
            path=resolved,
            sheets=tuple(sheets),
            workbook_theme=getattr(workbook, "loaded_theme", None),
        )
    finally:
        workbook.close()

_CACHED_INSPECT = None


def inspect_workbook_cached(path: Path) -> WorkbookInspection:
    """Streamlit cache wrapper keyed by workbook bytes, while keeping the pure parser reusable."""
    global _CACHED_INSPECT
    from metaphor_agreement_studio.common.cache_keys import workbook_content_cache_key

    if _CACHED_INSPECT is None:
        import streamlit as st

        @st.cache_data(show_spinner=False)
        def _cached(content_key: str, _path: str) -> WorkbookInspection:
            del content_key
            return inspect_workbook(Path(_path))

        _CACHED_INSPECT = _cached
    return _CACHED_INSPECT(workbook_content_cache_key(path), str(Path(path).resolve()))
